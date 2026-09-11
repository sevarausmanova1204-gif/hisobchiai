from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import config

_HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True)


def _style_header(ws, columns: list[str]) -> None:
    for col_idx, title in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = max(len(title) + 4, 14)


def _safe_amount(value) -> float:
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def _compute_summary(transactions: list[dict]) -> dict:
    total_kirim = 0.0
    total_chiqim = 0.0
    by_category: dict[str, float] = {}

    for row in transactions:
        summa = _safe_amount(row.get("Summa"))
        if row.get("Turi") == "Kirim":
            total_kirim += summa
        else:
            total_chiqim += summa
            kategoriya = row.get("Kategoriya", "")
            by_category[kategoriya] = by_category.get(kategoriya, 0.0) + summa

    return {
        "total_kirim": total_kirim,
        "total_chiqim": total_chiqim,
        "by_category": by_category,
    }


def _format_amount(amount: float) -> str:
    return f"{amount:,.0f}".replace(",", " ")


def build_text_report(transactions: list[dict], max_rows: int = 30) -> str:
    if not transactions:
        return "Hozircha hech qanday yozuv mavjud emas."

    summary = _compute_summary(transactions)
    balans = summary["total_kirim"] - summary["total_chiqim"]

    lines = ["📊 <b>Hisobot</b>\n"]

    recent = transactions[-max_rows:]
    for row in recent:
        emoji = "🟢" if row.get("Turi") == "Kirim" else "🔴"
        summa = _format_amount(_safe_amount(row.get("Summa")))
        valyuta = row.get("Valyuta", "")
        lines.append(
            f"{emoji} {row.get('Sana', '')} — {summa} {valyuta} "
            f"| {row.get('Kategoriya', '')} | {row.get('Tavsif', '')}"
        )

    if len(transactions) > max_rows:
        lines.append(f"\n… va yana {len(transactions) - max_rows} ta yozuv (to'liq ro'yxat Excel faylida)")

    lines.append("\n<b>Xulosa</b>")
    lines.append(f"🟢 Jami kirim: {_format_amount(summary['total_kirim'])}")
    lines.append(f"🔴 Jami chiqim: {_format_amount(summary['total_chiqim'])}")
    lines.append(f"💼 Balans: {_format_amount(balans)}")

    if summary["by_category"]:
        lines.append("\n<b>Kategoriya bo'yicha chiqimlar</b>")
        for cat, amount in sorted(summary["by_category"].items(), key=lambda x: x[1], reverse=True):
            lines.append(f"  • {cat}: {_format_amount(amount)}")

    return "\n".join(lines)


def build_excel(transactions: list[dict], filepath: str) -> str:
    wb = Workbook()

    ws = wb.active
    ws.title = "Tranzaksiyalar"
    _style_header(ws, config.SHEET_HEADERS)

    for row_idx, row in enumerate(transactions, start=2):
        ws.cell(row=row_idx, column=1, value=row.get("Sana", ""))
        ws.cell(row=row_idx, column=2, value=row.get("Turi", ""))
        ws.cell(row=row_idx, column=3, value=_safe_amount(row.get("Summa")))
        ws.cell(row=row_idx, column=4, value=row.get("Valyuta", ""))
        ws.cell(row=row_idx, column=5, value=row.get("Kategoriya", ""))
        ws.cell(row=row_idx, column=6, value=row.get("Tavsif", ""))
        ws.cell(row=row_idx, column=7, value=row.get("Foydalanuvchi", ""))

    summary = _compute_summary(transactions)
    total_kirim = summary["total_kirim"]
    total_chiqim = summary["total_chiqim"]
    by_category = summary["by_category"]

    summary_ws = wb.create_sheet("Xulosa")
    _style_header(summary_ws, ["Ko'rsatkich", "Qiymat"])
    summary_ws.cell(row=2, column=1, value="Jami kirim")
    summary_ws.cell(row=2, column=2, value=total_kirim)
    summary_ws.cell(row=3, column=1, value="Jami chiqim")
    summary_ws.cell(row=3, column=2, value=total_chiqim)
    summary_ws.cell(row=4, column=1, value="Balans")
    summary_ws.cell(row=4, column=2, value=total_kirim - total_chiqim)

    start_row = 6
    summary_ws.cell(row=start_row, column=1, value="Kategoriya bo'yicha chiqimlar").font = Font(bold=True)
    for i, (cat, amount) in enumerate(
        sorted(by_category.items(), key=lambda x: x[1], reverse=True), start=start_row + 1
    ):
        summary_ws.cell(row=i, column=1, value=cat)
        summary_ws.cell(row=i, column=2, value=amount)

    summary_ws.column_dimensions["A"].width = 30
    summary_ws.column_dimensions["B"].width = 18

    wb.save(filepath)
    return filepath
