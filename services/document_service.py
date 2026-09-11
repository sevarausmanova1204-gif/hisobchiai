import os

import pdfplumber
from docx import Document
from openpyxl import load_workbook

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".xlsx", ".csv"}


def extract_text(filepath: str) -> str:
    ext = os.path.splitext(filepath)[1].lower()

    if ext == ".pdf":
        return _extract_pdf(filepath)
    if ext == ".docx":
        return _extract_docx(filepath)
    if ext in (".txt", ".csv"):
        return _extract_plain_text(filepath)
    if ext == ".xlsx":
        return _extract_xlsx(filepath)

    raise ValueError(f"Qo'llab-quvvatlanmaydigan fayl turi: {ext}")


def _extract_pdf(filepath: str) -> str:
    chunks = []
    with pdfplumber.open(filepath) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                chunks.append(text)
    return "\n".join(chunks)


def _extract_docx(filepath: str) -> str:
    doc = Document(filepath)
    parts = [p.text for p in doc.paragraphs if p.text]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _extract_plain_text(filepath: str) -> str:
    with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


def _extract_xlsx(filepath: str) -> str:
    wb = load_workbook(filepath, data_only=True)
    lines = []
    for sheet in wb.worksheets:
        lines.append(f"--- {sheet.title} ---")
        for row in sheet.iter_rows(values_only=True):
            values = [str(v) for v in row if v is not None]
            if values:
                lines.append(" | ".join(values))
    return "\n".join(lines)
